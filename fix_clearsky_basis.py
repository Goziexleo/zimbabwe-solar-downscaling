"""Disclose the clear-sky time-base mismatch, which makes a QC claim vacuous.

The target is a 24-hour mean: ERA5 ssrd daily accumulations divided by 86,400 s.
The clear-sky denominator is NOT. compute_finegrid_clearsky_ghi.py averages the
PVLIB Ineichen ceiling over thirteen samples spanning 06:00 to 18:00 and divides
by thirteen, giving a DAYTIME mean of 502.7 W/m2 against a target whose mean is
237.3. The two differ by a factor of almost exactly two, so the quantity called a
clear-sky index has a mean of 0.476 where a consistent ratio would give 0.944.

This does not affect accuracy. The denominator is a per-cell, per-month constant,
so the mapping is a fixed monotonic rescaling, csi_to_ghi inverts it exactly, and
every metric in Chapter 4 is computed in GHI space after conversion.

It does invalidate one claim. Section 3.4.3 reports that the clear-sky index
reaches a maximum of 0.6323, that the 1.1 clip therefore never binds, and that
"the number of flagged ERA5 values is exactly zero" - presented as evidence that
ERA5 never exceeds its own clear-sky ceiling. On a consistent 24-hour basis that
maximum is 1.265, which exceeds the clip. The test cannot bind as constructed, so
its passing says nothing about ERA5, and the sentence is corrected to say so.

Re-deriving the target on a consistent basis would change the clipping and so
requires a full retrain. It belongs with the domain fix in RUNBOOK_DOMAIN_FIX.md
rather than being attempted piecemeal.
"""

import os
import shutil
import sys

import docx

CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

OLD = ("The constraint does not bind anywhere in the record: across 1,794,312 "
       "training and 966,168 validation values the clear-sky index reaches a "
       "maximum of 0.6323, so the number of flagged ERA5 values is exactly zero.")
NEW = ("The constraint does not bind anywhere in the record: across 1,794,312 "
       "training and 966,168 validation values the clear-sky index reaches a "
       "maximum of 0.6323, so the number of flagged ERA5 values is exactly zero. "
       "That result is an artefact of the two quantities being on different time "
       "bases, and is reported here rather than left to be discovered. The target "
       "is a 24-hour mean, ERA5 daily accumulations divided by 86,400 seconds, "
       "whereas the clear-sky ceiling is averaged over thirteen samples spanning "
       "06:00 to 18:00 and is therefore a daytime mean: 502.7 W/m² against a "
       "target mean of 237.3. The ratio of the two conventions is almost exactly "
       "two, which is why the index has a mean of 0.476 where a consistent ratio "
       "would give 0.944, and why its maximum of 0.6323 becomes 1.265 on a "
       "consistent basis and would exceed the 1.1 bound. The clip as constructed "
       "cannot bind, so its passing is not evidence about ERA5. Accuracy is "
       "unaffected, because the denominator is a per-cell, per-month constant and "
       "the conversion back to irradiance inverts it exactly, and every metric in "
       "Chapter 4 is computed in irradiance units after that conversion. Correcting "
       "the basis would change which values the clip truncates and so requires the "
       "models to be refitted; it is recorded as outstanding in Section 4.9.")

OLD2 = "producing a dimensionless clear-sky index (CSI), bounded as described below,"
NEW2 = ("producing a dimensionless ratio to the clear-sky ceiling, referred to "
        "throughout as the clear-sky index although Section 3.4.3 records that its "
        "denominator is on a daytime rather than a 24-hour basis,")


def is_field(r):
    x = r._r.xml
    return ("fldChar" in x) or ("instrText" in x)


def apply(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return True
    runs = par.runs
    full = "".join(r.text for r in runs)
    at = full.find(old)
    if at < 0:
        return False
    starts, pos = [], 0
    for r in runs:
        starts.append(pos)
        pos += len(r.text)
    f = max(i for i, st in enumerate(starts) if st <= at)
    l = max(i for i, st in enumerate(starts) if st < at + len(old))
    span = runs[f:l + 1]
    if any(is_field(r) for r in span):
        return False
    acc = "".join(r.text for r in span)
    k = acc.find(old)
    span[0].text = acc[:k] + new + acc[k + len(old):]
    for r in span[1:]:
        r.text = ""
    return True


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2
    d = docx.Document(DOC)
    x = d.element.xml
    before = (x.count("ZOTERO_ITEM"), x.count('w:fldCharType="begin"'))
    print("before:", before)

    done = 0
    for old, new in ((OLD, NEW), (OLD2, NEW2)):
        par = next((p for p in d.paragraphs if old in p.text), None)
        if par is None:
            print("  absent or already applied: %s" % old[:54])
            continue
        if new in par.text:
            print("  already applied")
            continue
        if apply(par, old, new):
            done += 1
            print("  applied: %s" % old[:54])
        else:
            print("  SPANS A FIELD, left: %s" % old[:54])

    if not done:
        print("nothing to do")
        return 0
    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_clearsky.docx"))
    d.save(DOC)
    y = docx.Document(DOC).element.xml
    after = (y.count("ZOTERO_ITEM"), y.count('w:fldCharType="begin"'))
    if after != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_clearsky.docx"), DOC)
        print("CITATION FIELDS CHANGED, restored")
        return 1
    print("%d applied; citation fields intact" % done)
    return 0


if __name__ == "__main__":
    sys.exit(main())
