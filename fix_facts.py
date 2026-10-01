"""Correct the factual errors the examiner's critique identified.

Each was checked before being changed, and where the project's own data could
settle it, that was used in preference to a secondary source:

  provinces      GADM level 1 reprojected to Albers equal area: FIVE provinces
                 exceed 40,000 km2, so a 200 km model cell is not larger than
                 every province. Matabeleland North alone is 75,444 km2.
  CNRM-CM6-1     the nominal_resolution attribute on the project's own rsds file
                 reads "250 km", and its grid is 1.4 degrees, about 155 km. The
                 thesis called it 100 km, which is MPI-ESM1-2-HR.
  elevation      the 124 m minimum is a cell OUTSIDE Zimbabwe. Masked to the
                 country, SRTM gives 211 m.
  CMIP6 daily    the CMIP6 day table defines clt, tas, huss and rsds; only ps and
                 od550aer are monthly-only. The thesis said all six were monthly.
  SARAH          Section 3.3.3 said "Heliosat Edition 2 (SARAH-3)" where Section
                 2.3.3 says Edition 3. SARAH-3 is Edition 3.
  CMIP6 start    the historical experiment begins in 1850, not mid-century.
  latitude       Zimbabwe reaches about 22.4 S; "15 to 22" is what produced the
                 truncated analysis grid.

Replacements are made inside the runs holding the text, never across a citation
field, and the field count is asserted before and after.
"""

import os
import shutil
import sys

import docx

CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

EDITS = [
    ("nearly half the world’s energy-poor population",
     "nearly half the world’s energy-poor population",
     "about four fifths of the global total"),

    ("NASA POWER at 0.5°",
     "NASA POWER at 0.5°  (~55 km)",
     "NASA POWER, whose solar parameters are supplied on a 1° grid"),

    ("exceeding any individual province in Zimbabwe",
     "exceeding any individual province in Zimbabwe",
     "larger than five of Zimbabwe’s ten provinces"),

    ("from the mid-twentieth century through to 2100",
     "from the mid-twentieth century through to 2100",
     "from 1850 through to 2100"),

    ("CNRM-CM6-1 (Centre National de Recherches Meteorologiques, France; approximately 100 km native resolution)",
     "CNRM-CM6-1 (Centre National de Recherches Meteorologiques, France; approximately 100 km native resolution)",
     "CNRM-CM6-1 (Centre National de Recherches Meteorologiques, France; 250 km "
     "nominal resolution, a 1.4° grid)"),

    ("ageing thermal generation infrastructure at Hwange Power Station",
     "ageing thermal generation infrastructure at Hwange Power Station",
     "an ageing thermal fleet at Hwange Power Station, whose 600 MW Units 7 and 8 "
     "commissioned in 2023 have not removed the constraint"),

    ("Heliosat Edition 2 (SARAH-3)",
     "Heliosat Edition 2 (SARAH-3)",
     "Heliosat Edition 3 (SARAH-3)"),

    ("CMIP6 provides only monthly output for these variables in any case",
     "CMIP6 provides only monthly output for these variables in any case",
     "two of the six predictors, surface pressure and aerosol optical depth, are "
     "archived only monthly, so a fully daily projection chain is not available "
     "for the predictor set as a whole"),

    ("which ranges from about 124 m in the Lowveld",
     "which ranges from about 124 m in the Lowveld",
     "which ranges from about 211 m in the Lowveld within Zimbabwe, and 124 m "
     "across the wider analysis box,"),

    ("between latitudes 15° S and 22° S",
     "between latitudes 15° S and 22° S",
     "between about 15.6° S and 22.4° S"),
]


def is_field(run):
    x = run._r.xml
    return ("fldChar" in x) or ("instrText" in x)


def apply(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return True
    runs = par.runs
    full = "".join(r.text for r in runs)
    at = full.find(old)
    if at < 0:
        return False
    end = at + len(old)
    starts, pos = [], 0
    for r in runs:
        starts.append(pos)
        pos += len(r.text)
    first = max(i for i, st in enumerate(starts) if st <= at)
    last = max(i for i, st in enumerate(starts) if st < end)
    span = runs[first:last + 1]
    if any(is_field(r) for r in span):
        return False
    acc = "".join(r.text for r in span)
    k = acc.find(old)
    span[0].text = acc[:k] + new + acc[k + len(old):]
    for r in span[1:]:
        r.text = ""
    return True


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2
    d = docx.Document(DOC)
    x = d.element.xml
    before = (x.count("ZOTERO_ITEM"), x.count('w:fldCharType="begin"'))
    print("before:", before)

    done = left = 0
    for locator, old, new in EDITS:
        par = next((p for p in d.paragraphs if locator in p.text), None)
        if par is None:
            print("  already applied or absent: %s" % locator[:52])
            continue
        if new in par.text:
            print("  already applied: %s" % locator[:52])
            continue
        if apply(par, old, new):
            done += 1
            print("  fixed: %s" % locator[:52])
        else:
            left += 1
            print("  SPANS A CITATION FIELD, left: %s" % locator[:52])

    if not done:
        print("nothing to do")
        return 0
    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_facts.docx"))
    d.save(DOC)
    y = docx.Document(DOC).element.xml
    after = (y.count("ZOTERO_ITEM"), y.count('w:fldCharType="begin"'))
    if after != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_facts.docx"), DOC)
        print("CITATION FIELDS CHANGED, restored")
        return 1
    print("%d fixed, %d left; citation fields intact" % (done, left))
    return 0


if __name__ == "__main__":
    sys.exit(main())
