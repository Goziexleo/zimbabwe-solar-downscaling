"""Stop citing Harilal et al. (2022) for the resolution of the CORDEX archive.

Section 2.x reads "...at 0.22 degrees and 0.44 degrees resolution over Africa
(Harilal et al., 2022)". That paper is EnhancedSD, a NeurIPS workshop paper on
downscaling solar irradiance from climate model projections with machine
learning. It is a statistical-downscaling paper and says nothing about the
resolution of a dynamical archive, so the attribution is wrong.

The citation is a live Zotero field. The edit therefore touches only the plain
text run immediately before it, closing the CORDEX sentence and opening a new
clause that the citation does support. The resolution figures are left without a
citation and listed in the handover as needing a CORDEX source added in Zotero,
which is work only the author can do.

    python fix_cordex_misattribution.py
"""

import os
import shutil
import sys

import docx

CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

OLD = " (~50 km) resolution over Africa "
NEW = (" (~50 km) resolution over Africa. The same downscaling problem has been "
       "approached statistically for solar irradiance with machine learning ")


def is_field(r):
    x = r._r.xml
    return ("fldChar" in x) or ("instrText" in x)


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2
    d = docx.Document(DOC)
    before = (d.element.xml.count("ZOTERO_ITEM"),
              d.element.xml.count('w:fldCharType="begin"'))

    par = next((p for p in d.paragraphs
                if "CORDEX framework has produced a substantial archive" in p.text), None)
    if par is None:
        print("paragraph not found")
        return 1
    if NEW.strip() in par.text:
        print("already applied")
        return 0

    hits = [r for r in par.runs if r.text == OLD and not is_field(r)]
    if len(hits) != 1:
        print("expected exactly one plain-text run %r, found %d" % (OLD, len(hits)))
        return 1
    hits[0].text = NEW

    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_cordex.docx"))
    d.save(DOC)
    after_xml = docx.Document(DOC).element.xml
    after = (after_xml.count("ZOTERO_ITEM"), after_xml.count('w:fldCharType="begin"'))
    if after != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_cordex.docx"), DOC)
        print("CITATION FIELDS CHANGED %s -> %s, restored" % (before, after))
        return 1
    print("fixed; citation fields intact %s" % (after,))
    return 0


if __name__ == "__main__":
    sys.exit(main())
