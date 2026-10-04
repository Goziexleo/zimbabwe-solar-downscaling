"""Clear the wording and attribution defects the fourth critique still lists.

Four of them, all plain text sitting beside citation fields rather than inside
them, so each edit touches only runs the field guard allows:

  - "Lin et al., 2023 Applied a U-Net architecture" - a sentence beginning with
    a citation that was never turned into narrative form
  - "...aerosol optical depth. (Eyring et al., 2016). note that inter-model
    spread..." - a stray full stop before the field and a lower-case fragment
    after it, left from an edit that moved the citation
  - Damiani et al. (2024) "demonstrated that higher spatial fidelity ...
    produces more accurate PV yield estimates" - the second of two overclaims;
    the first was softened in an earlier round and this one was missed
  - "Section 3.8.4 originally deployed the Random Forest" - revision history in
    a sentence whose point stands without it

    python fix_v4_wording.py
"""

import os
import shutil
import sys

import docx

CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

EDITS = [
    # The citation itself is a Zotero field rendering "Lin et al., 2023", so only
    # the text after it can be touched here. Lower-casing the verb removes the
    # sentence-initial capital mid-sentence; putting the year in brackets needs
    # the field switched to narrative mode in Word, which is listed in the
    # handover as an author action.
    ("Applied a U-Net architecture to downscale CMIP6 variables",
     "applied a U-Net architecture to downscale CMIP6 variables"),

    # The field sits between these two fragments; both are plain text.
    ("underestimation of low-level cloud fraction and aerosol optical depth. ",
     "underestimation of low-level cloud fraction and aerosol optical depth "),
    (". note that inter-model spread in rsds over Africa remains substantial",
     ". Inter-model spread in rsds over Africa remains substantial"),

    ("demonstrated that higher spatial fidelity in downscaled solar radiation fields "
     "produces more accurate PV yield estimates",
     "report that higher spatial fidelity in downscaled solar radiation fields is "
     "associated with closer agreement in PV yield estimates"),

    ("This matters because Section 3.8.4 originally deployed the Random Forest on a "
     "composite criterion combining systematic offset with spatial error structure.",
     "This matters because Section 3.8.4 weighs a composite criterion combining "
     "systematic offset with spatial error structure, under which the Random Forest "
     "would be preferred."),
]


def is_field(r):
    x = r._r.xml
    return ("fldChar" in x) or ("instrText" in x)


def apply(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return True
    runs = par.runs
    full = "".join(r.text for r in runs)
    at = full.find(old)
    if at < 0:
        return False
    starts, pos = [], 0
    for r in runs:
        starts.append(pos)
        pos += len(r.text)
    f = max(i for i, st in enumerate(starts) if st <= at)
    l = max(i for i, st in enumerate(starts) if st < at + len(old))
    span = runs[f:l + 1]
    if any(is_field(r) for r in span):
        return False
    acc = "".join(r.text for r in span)
    k = acc.find(old)
    span[0].text = acc[:k] + new + acc[k + len(old):]
    for r in span[1:]:
        r.text = ""
    return True


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2
    d = docx.Document(DOC)
    before = (d.element.xml.count("ZOTERO_ITEM"),
              d.element.xml.count('w:fldCharType="begin"'))
    done = left = 0
    for old, new in EDITS:
        par = next((p for p in d.paragraphs if old in p.text), None)
        if par is None:
            print("  already applied or absent: %s" % old[:52])
            continue
        if apply(par, old, new):
            done += 1
            print("  fixed: %s" % old[:52])
        else:
            left += 1
            print("  SPANS A FIELD, skipped: %s" % old[:52])
    if not done:
        return 0 if left == 0 else 1
    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_v4wording.docx"))
    d.save(DOC)
    ax = docx.Document(DOC).element.xml
    after = (ax.count("ZOTERO_ITEM"), ax.count('w:fldCharType="begin"'))
    if after != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_v4wording.docx"), DOC)
        print("CITATION FIELDS CHANGED %s -> %s, restored" % (before, after))
        return 1
    print("%d fixed, %d outstanding; citation fields intact %s" % (done, left, after))
    return 0


if __name__ == "__main__":
    sys.exit(main())
