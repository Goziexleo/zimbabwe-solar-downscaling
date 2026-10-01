"""Qualify Chapter 1's justification for machine learning, which Section 4.2.1 tests.

Chapter 1 argued that classical statistical downscaling is linear, that the
predictor-predictand relationship for irradiance is not, and that machine
learning is therefore warranted. The first two claims are about the atmosphere
and stand. The inference does not: a per-cell ordinary least squares regression
on the same predictors reaches 8.75 W m-2 against the deployed XGBoost's 9.14,
indistinguishable from it on a paired year-block bootstrap and established
better than the other three architectures.

The paragraphs are not rewritten to claim the opposite. They are changed from
asserting the premise to posing it as the question the study answers, which is
what the evidence now supports and what makes Chapter 1 agree with Chapter 4.

Edits are applied run by run, and the citation-field count is checked before and
after, because a Zotero field occupies several runs and flattening a paragraph
destroys it - which has happened in this project before.
"""

import os
import shutil
import sys

import docx

CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

ADD_57 = (" These are claims about the atmosphere, not about any particular "
          "downscaling task. Whether the nonlinearity they describe is recoverable "
          "depends on what the model is trained against, so this study treats the "
          "premise as a question rather than an assumption: Section 4.2.1 fits a "
          "per-cell linear regression on the same predictors and reports what the "
          "nonlinear models are worth against it.")

OLD_59 = "Machine learning offers a compelling alternative to classical statistical downscaling by relaxing the linearity constraint."
NEW_59 = "Machine learning offers an alternative to classical statistical downscaling by relaxing the linearity constraint."

PARA_AFTER_59 = (
    "Whether that relaxation buys anything is an empirical question, and on the "
    "target used here the answer turns out to be largely negative. Section 4.2.1 "
    "reports that a per-cell linear regression on the same predictors is "
    "indistinguishable from the best of the four architectures and better than the "
    "other three. Section 4.5 gives the reason: the training target retains almost "
    "none of its variance below the resolution of its own source, so it is close to "
    "a linear function of the coarse predictors and offers little nonlinear "
    "structure to learn. The architectures are still compared, and the comparison "
    "is the substance of Chapter 4, but the justification above should be read as "
    "the hypothesis this study tests rather than as a settled premise.")


def fields(d):
    x = d.element.xml
    return (x.count("ZOTERO_ITEM"), x.count('w:fldCharType="begin"'),
            x.count('w:fldCharType="end"'))


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2

    d = docx.Document(DOC)
    before = fields(d)
    print("before:", before)

    p57 = next(p for p in d.paragraphs
               if p.text.strip().startswith("A fundamental limitation of classical"))
    if ADD_57.strip() in p57.text:
        print("57 already qualified")
    else:
        p57.runs[-1].text = p57.runs[-1].text + ADD_57
        print("57 qualified (%d runs, untouched count)" % len(p57.runs))

    p59 = next(p for p in d.paragraphs
               if p.text.strip().startswith("Machine learning offers"))
    hit = False
    for r in p59.runs:
        if OLD_59 in r.text:
            r.text = r.text.replace(OLD_59, NEW_59)
            hit = True
            break
    print("59 opening softened" if hit else "59 opening already soft")

    # New paragraph straight after 59, matching its style.
    if not any(p.text.strip().startswith("Whether that relaxation buys")
               for p in d.paragraphs):
        new = d.add_paragraph(PARA_AFTER_59, style=p59.style)
        p59._p.addnext(new._p)
        pf, src = new.paragraph_format, p59.paragraph_format
        pf.alignment = p59.alignment
        pf.space_before, pf.space_after = src.space_before, src.space_after
        pf.line_spacing, pf.first_line_indent = src.line_spacing, src.first_line_indent
        for r in new.runs:
            if p59.runs:
                r.font.name = p59.runs[0].font.name
                r.font.size = p59.runs[0].font.size
        print("added the forward-reference paragraph")

    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_reframe.docx"))
    d.save(DOC)

    after = fields(docx.Document(DOC))
    print("after: ", after)
    if after != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_reframe.docx"), DOC)
        print("CITATION FIELDS CHANGED, restored from backup")
        return 1
    print("citation fields intact; saved")
    return 0


if __name__ == "__main__":
    sys.exit(main())
