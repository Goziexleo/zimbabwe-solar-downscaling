"""Give Table 3.2 the final hyperparameter values its own text promises.

Section 3.6.7 says twice that Table 3.2 reports the selected values - "Table 3.2
reports the honestly selected values" and "Final hyperparameter values are
reported in Table 3.2" - and the table reports no values at all. It lists the
hyperparameters, the range searched over each and the selection criterion, and
stops. A reader cannot reproduce any of the four models from it.

A "Deployed value" column is added, parsed from the training scripts themselves
so the table cannot drift from what the code does. Two sentences in Section 3.6.7
are corrected at the same time:

  - the claim that the table holds the "honestly selected" values, which is wrong
    for the two pixel-wise models: the same paragraph already says their defaults
    were retained in preference to the search result
  - the neural inner-split figures, which come from a search run before the
    clear-sky rebuild and are not the deployed models' own scores

The cost of retaining the defaults is then stated as a number, from the search
re-run with the deployed configuration included as an extra point.

    python fix_table_3_2.py
"""

import os
import re
import shutil
import sys

import docx
import pandas as pd
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = os.path.dirname(os.path.abspath(__file__))
EVAL = os.path.join(ROOT, "data/processed/evaluation")
CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

CAPTION = "Table 3.2. Hyperparameter optimisation specifications"
HEADER = "Deployed value"


def default_of(script, var, name):
    """Read a training script's env-overridable default, so the table matches code."""
    src = open(os.path.join(ROOT, script)).read()
    m = re.search(r'%s\s*=\s*\w+\(os\.environ\.get\(\s*"%s"\s*,\s*"([^"]+)"' % (var, name), src)
    if m is None:
        m = re.search(r'os\.environ\.get\(\s*"%s"\s*,\s*"([^"]+)"' % name, src)
    assert m is not None, "could not read %s from %s" % (name, script)
    return m.group(1)


def deployed_values():
    rf = "n_estimators %s; max_features %s; min_samples_leaf %s" % (
        default_of("train_pixelwise_rf.py", "N_ESTIMATORS", "RF_N_ESTIMATORS"),
        default_of("train_pixelwise_rf.py", "_max_features_raw", "RF_MAX_FEATURES"),
        default_of("train_pixelwise_rf.py", "MIN_SAMPLES_LEAF", "RF_MIN_SAMPLES_LEAF"))
    xgb = "max_depth %s; eta %s; subsample %s; min_child_weight %s; %s rounds" % (
        default_of("train_pixelwise_xgb.py", "MAX_DEPTH", "XGB_MAX_DEPTH"),
        default_of("train_pixelwise_xgb.py", "ETA", "XGB_ETA"),
        default_of("train_pixelwise_xgb.py", "SUBSAMPLE", "XGB_SUBSAMPLE"),
        default_of("train_pixelwise_xgb.py", "MIN_CHILD_WEIGHT", "XGB_MIN_CHILD_WEIGHT"),
        default_of("train_pixelwise_xgb.py", "N_ESTIMATORS", "XGB_N_ESTIMATORS"))
    cnn = "learning rate %s; loss weighting %s" % (
        default_of("train_cnn_downscaler.py", "LEARNING_RATE", "CNN_LEARNING_RATE"),
        default_of("train_cnn_downscaler.py", "LAMBDA_GP", "CNN_LAMBDA_GP"))
    def dec(v):
        """0.001 rather than 1e-3: the table is read, not parsed."""
        f = float(v)
        return ("%.10f" % f).rstrip("0").rstrip(".") if f else "0"

    unet = "learning rate %s; dropout %s (fixed, not searched)" % (
        dec(default_of("train_unet_downscaler.py", "LEARNING_RATE", "UNET_LEARNING_RATE")),
        dec(default_of("unet_model.py", "DROPOUT_P", "UNET_DROPOUT")))
    return {"Random Forest": rf, "XGBoost": xgb, "CNN": cnn, "U-Net": unet}


def rank_line(model, csv):
    path = os.path.join(EVAL, csv)
    if not os.path.exists(path):
        return None
    d = pd.read_csv(path)
    dep = d[d.is_deployed]
    if dep.empty:
        return None
    r = dep.iloc[0]
    best = float(d.cv_rmse_csi.min())
    return (model, int(r["rank"]), len(d), float(r["cv_rmse_csi"]), best,
            100.0 * (float(r["cv_rmse_csi"]) - best) / best)


def add_column(tbl, values_by_first_cell):
    """Append a column, widening the grid, and fill it by the row's first cell."""
    grid = tbl._tbl.find(qn("w:tblGrid"))
    assert grid is not None, "table has no grid"
    cols = grid.findall(qn("w:gridCol"))
    width = int(cols[-1].get(qn("w:w")) or 1800)
    gc = OxmlElement("w:gridCol")
    gc.set(qn("w:w"), str(width))
    grid.append(gc)

    for i, row in enumerate(tbl.rows):
        tc = OxmlElement("w:tc")
        tcPr = OxmlElement("w:tcPr")
        cw = OxmlElement("w:tcW")
        cw.set(qn("w:w"), str(width))
        cw.set(qn("w:type"), "dxa")
        tcPr.append(cw)
        tc.append(tcPr)
        tc.append(OxmlElement("w:p"))
        row._tr.append(tc)
        cell = row.cells[-1]
        key = row.cells[0].text.strip()
        text = HEADER if i == 0 else values_by_first_cell.get(key, "—")
        run = cell.paragraphs[0].add_run(text)
        run.font.size = Pt(10)
        if i == 0:
            run.bold = True


def is_field(r):
    x = r._r.xml
    return ("fldChar" in x) or ("instrText" in x)


def apply(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return True
    runs = par.runs
    full = "".join(r.text for r in runs)
    at = full.find(old)
    if at < 0:
        return False
    starts, pos = [], 0
    for r in runs:
        starts.append(pos)
        pos += len(r.text)
    f = max(i for i, st in enumerate(starts) if st <= at)
    l = max(i for i, st in enumerate(starts) if st < at + len(old))
    span = runs[f:l + 1]
    if any(is_field(r) for r in span):
        return False
    acc = "".join(r.text for r in span)
    k = acc.find(old)
    span[0].text = acc[:k] + new + acc[k + len(old):]
    for r in span[1:]:
        r.text = ""
    return True


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2

    vals = deployed_values()
    for k, v in vals.items():
        print("  %-15s %s" % (k, v))

    ranks = [r for r in (rank_line("XGBoost", "hpo_pixelwise_xgb.csv"),
                         rank_line("Random Forest", "hpo_pixelwise_rf.csv")) if r]
    if not ranks:
        print("no search results found; run hpo_pixelwise.py first")
        return 1

    d = docx.Document(DOC)
    before = (d.element.xml.count("ZOTERO_ITEM"),
              d.element.xml.count('w:fldCharType="begin"'))

    # The table sits immediately before its caption.
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    items = []
    for ch in d.element.body:
        if ch.tag.endswith("}p"):
            items.append(Paragraph(ch, d))
        elif ch.tag.endswith("}tbl"):
            items.append(Table(ch, d))
    caps = [n for n, it in enumerate(items)
            if isinstance(it, Paragraph) and it.text.strip().startswith(CAPTION)]
    assert caps, "Table 3.2 caption not found"
    tbl = None
    for cap in caps:
        for k in range(cap - 1, max(cap - 4, -1), -1):
            if isinstance(items[k], Table):
                tbl = items[k]
                break
        if tbl is not None:
            break
    assert tbl is not None, (
        "no table precedes any of the %d Table 3.2 captions" % len(caps))

    changed = False
    if HEADER in tbl.rows[0].cells[-1].text:
        print("  column already present")
    else:
        add_column(tbl, vals)
        changed = True
        print("  added the %r column (%d columns now)" % (HEADER, len(tbl.columns)))

    sentences = []
    for model, rk, n, cv, best, pct in ranks:
        sentences.append("the deployed %s ranks %d of %d at %.5f against %.5f for the "
                         "best, %.1f per cent higher" % (model, rk, n, cv, best, pct))
    cost = ("The cost of retaining them is now measured rather than asserted. With each "
            "deployed configuration added to its own grid as an extra point, %s. The "
            "XGBoost values for eta, subsample and minimum child weight all lie outside "
            "the grid of Table 3.2, so the search could never have returned that "
            "configuration; the Random Forest values do lie inside it."
            % ", and ".join(sentences))

    EDITS = [
        ("Table 3.2 reports the honestly selected values, and Section 4.2 quantifies "
         "what the change was worth on the withheld record.",
         "Table 3.2 reports the deployed value of every hyperparameter beside the range "
         "searched over it, and Section 4.2 quantifies what the change was worth on the "
         "withheld record. One caveat attaches to the two figures just quoted: they come "
         "from a search run before the clear-sky target was rebuilt, so they are not the "
         "inner-split scores of the models actually deployed, which are 0.001245 for the "
         "CNN and 0.120408 for the U-Net. The search was not repeated after the rebuild. "
         "What it established was the ordering, and the learning rate of 0.001 it selected "
         "is the one both deployed models carry."),

        ("The default configurations were therefore retained for the Random Forest and "
         "XGBoost, while the two neural architectures adopted their search-selected "
         "values.",
         "The default configurations were therefore retained for the Random Forest and "
         "XGBoost, while the two neural architectures adopted their search-selected "
         "values. " + cost),
    ]

    done = left = 0
    for old, new in EDITS:
        if any(new[-70:] in p.text for p in d.paragraphs):
            print("  already applied: %s" % old[:46])
            continue
        par = next((p for p in d.paragraphs if old in p.text), None)
        if par is None:
            print("  NOT FOUND: %s" % old[:46])
            left += 1
            continue
        if apply(par, old, new):
            done += 1
            changed = True
            print("  fixed: %s" % old[:46])
        else:
            left += 1
            print("  SPANS A FIELD, skipped: %s" % old[:46])

    if not changed:
        return 0 if left == 0 else 1
    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_table32.docx"))
    d.save(DOC)
    ax = docx.Document(DOC).element.xml
    after = (ax.count("ZOTERO_ITEM"), ax.count('w:fldCharType="begin"'))
    if after != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_table32.docx"), DOC)
        print("CITATION FIELDS CHANGED %s -> %s, restored" % (before, after))
        return 1
    print("%d sentences fixed, %d outstanding; citation fields intact %s"
          % (done, left, after))
    print("run normalise_tables.py afterwards: the new column needs the shared width")
    return 0


if __name__ == "__main__":
    sys.exit(main())
