"""Clear the text defects the third review found.

Most are single sentences. Two are substantive rather than cosmetic.

The ordering percentage in Section 3.8.4 quoted 96.7 per cent for XGBoost where
the screen now reports 99.1, and the preserved Chapter 5 paragraph still carried
71.7 against the screen's 71.5. Both are stale rather than a difference of basis.

"Categorical rather than a matter of degree" overstates the Random Forest's
failure: 71.5 per cent correct ordering is well above chance, so the failure is
large, not categorical. The word is changed because an examiner can check it.
"""

import os
import shutil
import sys

import docx

CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

EDITS = [
    ("against 96.7 per cent for XGBoost", "against 99.1 per cent for XGBoost"),
    ("and only 71.7 per cent of cells order the two pathways correctly",
     "and only 71.5 per cent of cells order the two pathways correctly"),
    ("the failure is categorical rather than a matter of degree",
     "the failure is large rather than marginal, though 71.5 per cent correct "
     "ordering is still well above chance"),
    ("derived from Meteosat retrievals using the HELIOSAT-2 algorithm",
     "derived from Meteosat retrievals using the Heliosat method with the "
     "SPECMAGIC clear-sky model"),
    ("referred to throughout as the clear-sky index although Section 3.4.3 records "
     "that its denominator is on a daytime rather than a 24-hour basis,",
     "referred to throughout as the clear-sky index,"),
    # Damiani evaluates PV output; it does not demonstrate that spatial fidelity
    # causes more accurate yield estimates.
    ("and that improved spatial fidelity translates into more accurate PV yield "
     "estimates.",
     "and report PV output computed from the downscaled fields."),
    ("Damiani et al. (2024) demonstrated that higher spatial fidelity in downscaled "
     "solar radiation fields produces more accurate PV yield estimates, which "
     "motivates",
     "Damiani et al. (2024) argue that higher spatial fidelity in downscaled solar "
     "radiation fields matters for PV yield estimation, which motivates"),
]


def is_field(r):
    x = r._r.xml
    return ("fldChar" in x) or ("instrText" in x)


def apply(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return True
    runs = par.runs
    full = "".join(r.text for r in runs)
    at = full.find(old)
    if at < 0:
        return False
    starts, pos = [], 0
    for r in runs:
        starts.append(pos)
        pos += len(r.text)
    f = max(i for i, st in enumerate(starts) if st <= at)
    l = max(i for i, st in enumerate(starts) if st < at + len(old))
    span = runs[f:l + 1]
    if any(is_field(r) for r in span):
        return False
    acc = "".join(r.text for r in span)
    k = acc.find(old)
    span[0].text = acc[:k] + new + acc[k + len(old):]
    for r in span[1:]:
        r.text = ""
    return True


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2
    d = docx.Document(DOC)
    x = d.element.xml
    before = (x.count("ZOTERO_ITEM"), x.count('w:fldCharType="begin"'))
    done = left = 0
    for old, new in EDITS:
        par = next((p for p in d.paragraphs if old in p.text), None)
        if par is None:
            print("  already applied or absent: %s" % old[:54])
            continue
        if apply(par, old, new):
            done += 1
            print("  fixed: %s" % old[:54])
        else:
            left += 1
            print("  SPANS A FIELD, left: %s" % old[:54])
    if not done:
        return 0
    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_v3.docx"))
    d.save(DOC)
    ax = docx.Document(DOC).element.xml
    after = (ax.count("ZOTERO_ITEM"), ax.count('w:fldCharType="begin"'))
    if after != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_v3.docx"), DOC)
        print("CITATION FIELDS CHANGED, restored")
        return 1
    print("%d fixed, %d left; citation fields intact" % (done, left))
    return 0


if __name__ == "__main__":
    sys.exit(main())
