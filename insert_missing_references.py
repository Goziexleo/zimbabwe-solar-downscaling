"""Add the works cited in the text but missing from the reference list.

Six sources are cited in the body and absent from the bibliography, and the
examiner has docked marks for it across four versions. The list is a Zotero
field, so the proper repair is the author importing reference_additions.ris and
refreshing; until that happens the document has six citations a reader cannot
resolve, which is the cheapest mark an examiner can take.

This inserts them as Bibliography-styled paragraphs in alphabetical position
inside the field's result, formatted to match the entries already there. A
Zotero refresh regenerates the whole block and replaces them, which is the
intended end state: nothing here has to be removed first.

Entries are built from reference_additions.ris, whose records for these six are
CrossRef-verified, and only surnames actually cited in the body and actually
missing from the list are added.

    python insert_missing_references.py [--dry-run]
"""

import argparse
import copy
import os
import re
import shutil
import sys

import docx

ROOT = os.path.dirname(os.path.abspath(__file__))
RIS = os.path.join(ROOT, "reference_additions.ris")
CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")
NDASH = "–"


def records():
    out = []
    for block in open(RIS).read().split("ER  -"):
        if "TY  -" not in block:
            continue
        f = {}
        for line in block.splitlines():
            m = re.match(r"^([A-Z][A-Z0-9])  - (.*)$", line.strip())
            if not m:
                continue
            k, v = m.group(1), m.group(2).strip()
            f.setdefault(k, []).append(v)
        if f:
            out.append(f)
    return out


def initials(name):
    """'Cohen, Jacob' -> 'Cohen, J.'; already-initialised names pass through."""
    if "," not in name:
        return name
    surname, given = name.split(",", 1)
    parts = [p for p in re.split(r"[\s.]+", given.strip()) if p]
    ini = "".join(p[0].upper() + "." for p in parts)
    return "%s, %s" % (surname.strip(), ini)


def authors(f):
    aus = f.get("AU", [])
    if len(aus) == 1 and aus[0].count(",") > 1:
        return aus[0]                      # already a formatted author string
    return ", ".join(initials(a) for a in aus)


def entry(f):
    bits = ["%s, %s." % (authors(f), f.get("PY", ["n.d."])[0])]
    title = f.get("TI", [""])[0].rstrip(".")
    bits.append(" %s." % title)
    jo = f.get("JO", [None])[0]
    if jo:
        vol = f.get("VL", [None])[0]
        sp, ep = f.get("SP", [None])[0], f.get("EP", [None])[0]
        tail = " %s" % jo
        if vol:
            tail += " %s" % vol
        if sp and ep:
            tail += ", %s%s%s" % (sp, NDASH, ep)
        elif sp:
            tail += ", %s" % sp
        bits.append(tail + ".")
    doi = f.get("DO", [None])[0]
    url = f.get("UR", [None])[0]
    if doi:
        bits.append(" https://doi.org/%s" % doi)
    elif url:
        bits.append(" %s" % url)
    return "".join(bits)


def sort_key(text):
    return re.sub(r"[^a-z]", "", text.split(",")[0].lower())


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2

    d = docx.Document(DOC)
    before = (d.element.xml.count("ZOTERO_ITEM"),
              d.element.xml.count('w:fldCharType="begin"'),
              d.element.xml.count('w:fldCharType="end"'))

    ps = list(d.paragraphs)
    ri = next(n for n, p in enumerate(ps) if p.text.strip().upper().startswith("REFERENCE"))
    bib_idx = [n for n, p in enumerate(ps) if n > ri
               and p.style.name == "Bibliography" and p.text.strip()]
    assert bib_idx, "no bibliography entries found"
    body = "\n".join(p.text for p in ps[:ri])
    listed = "\n".join(ps[n].text for n in bib_idx)

    pending = []
    for f in records():
        surname = authors(f).split(",")[0].strip()
        if not surname or surname not in body or surname in listed:
            continue
        pending.append((sort_key(entry(f)), entry(f), surname))
    pending.sort()

    if not pending:
        print("nothing to add: every cited work already appears in the list")
        return 0
    for _, text, surname in pending:
        print("  + %-14s %s" % (surname, text[:96]))
    if args.dry_run:
        print("\ndry run, nothing written")
        return 0

    template = ps[bib_idx[-1]]
    existing = [(sort_key(ps[n].text), n) for n in bib_idx]
    added = 0
    for key, text, _ in pending:
        after = None
        for k, n in existing:
            if k <= key:
                after = n
        target = ps[after if after is not None else bib_idx[0]]
        new_el = copy.deepcopy(template._p)
        # strip every run but one, then set the text, so the entry inherits the
        # paragraph formatting of the list it joins
        runs = new_el.findall(docx.oxml.ns.qn("w:r"))
        for r in runs[1:]:
            new_el.remove(r)
        para = docx.text.paragraph.Paragraph(new_el, target._parent)
        if not para.runs:
            para.add_run("")
        para.runs[0].text = text
        target._p.addnext(new_el)
        ps = list(d.paragraphs)
        bib_idx = [n for n, p in enumerate(ps) if p.style.name == "Bibliography"
                   and p.text.strip()]
        existing = [(sort_key(ps[n].text), n) for n in bib_idx]
        added += 1

    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_refs.docx"))
    d.save(DOC)
    x = docx.Document(DOC).element.xml
    after_counts = (x.count("ZOTERO_ITEM"), x.count('w:fldCharType="begin"'),
                    x.count('w:fldCharType="end"'))
    if after_counts != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_refs.docx"), DOC)
        print("FIELD COUNTS CHANGED %s -> %s, restored" % (before, after_counts))
        return 1
    n_now = sum(1 for p in docx.Document(DOC).paragraphs
                if p.style.name == "Bibliography" and p.text.strip())
    print("\nadded %d; the list now has %d entries; fields intact %s"
          % (added, n_now, after_counts))
    return 0


if __name__ == "__main__":
    sys.exit(main())
