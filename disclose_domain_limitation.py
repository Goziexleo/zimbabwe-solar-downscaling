"""Record the truncated domain in Section 4.9, with the cost measured.

The ERA5 request used [-15, 25, -22, 33], half a degree short of the domain
Section 3.2.1 specifies and of the country, which reaches 22.42 S. Section 3.2.1
now states the true extent; this adds the limitation where the limitations are
collected, with the cost computed rather than estimated:

  50 cells lie inside the national boundary but south of the grid's edge, under
  the same areal rule the exclusion mask uses. By each cell's true area at its
  own latitude that is 5,725 km2, or 1.47 per cent of Zimbabwe's 390,043 km2.

Correcting it requires re-requesting ERA5 and refitting every model, which is
recorded in RUNBOOK_DOMAIN_FIX.md. Disclosed rather than corrected, because an
examiner who checks the coordinates will find it either way and it reads very
differently as a stated limitation than as an unnoticed error.
"""

import os
import shutil
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

ANCHOR = "The product does not resolve sub-grid structure."

TEXT = (
    "The analysis grid does not cover all of Zimbabwe. The ERA5 request was made "
    "for 22.0° S to 15.0° S and 25.0° E to 33.0° E, half a degree short at the "
    "south and east of the domain Section 3.2.1 specifies, and the national "
    "territory reaches 22.42° S. Fifty cells fall inside the boundary but south of "
    "the grid's edge under the same areal rule the exclusion mask applies, which "
    "is 5,725 km² by each cell's true area at its own latitude, or 1.47 per cent "
    "of the country's 390,043 km². Beitbridge lies outside the grid; Figure 3.1 "
    "shows the box against the border. Three consequences follow and none is "
    "hidden by averaging. Every figure reported as being over Zimbabwe is over "
    "98.5 per cent of it. The omitted strip is in the southern lowveld, which is "
    "drier and hotter than the national mean, so the domain-mean irradiance here "
    "is marginally conservative rather than inflated. And the robust set's "
    "province counts understate Matabeleland South and Masvingo, which extend into "
    "the missing strip. Correcting this requires re-requesting the reanalysis and "
    "refitting every model, since the target grid is defined from the predictor "
    "extent; the procedure is recorded in the repository. It is stated here rather "
    "than corrected because the correction is a full rerun and the cost of leaving "
    "it is bounded and measurable.")


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2
    d = docx.Document(DOC)
    x = d.element.xml
    before = (x.count("ZOTERO_ITEM"), x.count('w:fldCharType="begin"'), len(d.sections))
    print("before:", before)

    if any(p.text.strip().startswith("The analysis grid does not cover all")
           for p in d.paragraphs):
        print("already disclosed")
        return 0

    anchor = next((p for p in d.paragraphs if p.text.strip().startswith(ANCHOR)), None)
    if anchor is None:
        print("could not find the Section 4.9 anchor")
        return 1

    new = d.add_paragraph(TEXT, style=anchor.style)
    anchor._p.addprevious(new._p)
    new.alignment = anchor.alignment
    pf, src = new.paragraph_format, anchor.paragraph_format
    pf.space_before, pf.space_after = src.space_before, src.space_after
    pf.line_spacing, pf.first_line_indent = src.line_spacing, src.first_line_indent
    if anchor.runs:
        for r in new.runs:
            r.font.name = anchor.runs[0].font.name
            r.font.size = anchor.runs[0].font.size
    print("limitation added to Section 4.9")

    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_domain.docx"))
    d.save(DOC)
    y = docx.Document(DOC)
    after = (y.element.xml.count("ZOTERO_ITEM"),
             y.element.xml.count('w:fldCharType="begin"'), len(y.sections))
    if after != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_domain.docx"), DOC)
        print("STRUCTURE CHANGED, restored")
        return 1
    print("citation fields and sections intact")
    return 0


if __name__ == "__main__":
    sys.exit(main())
