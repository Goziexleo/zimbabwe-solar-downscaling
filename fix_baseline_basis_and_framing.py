"""Put the interpolation-baseline sentences on the Zimbabwe basis and stop
calling that baseline the benchmark.

Two leftovers from before the basis alignment:

  3.6.5  says the validation period is used for "benchmarking against the naive
         interpolation baseline", which Section 3.7.3 then spends a paragraph
         denying: that baseline is a near-identity and is reported as a
         circularity diagnostic, not as a competitor.
  3.7.3  gives the baseline's error as 0.23 W/m2 and the climatology's as
         19.08 W/m2. The first matches neither basis (it is 0.2168 over
         Zimbabwe and 0.2420 over the full box) and the second is the full-box
         figure, while Table 4.1 and Section 4.6.1 both report the Zimbabwe
         basis. Both are moved to the Zimbabwe basis and the basis is named.

Numbers come from baselines.csv so they cannot drift from the computation.

    python fix_baseline_basis_and_framing.py
"""

import os
import shutil
import sys

import docx
import pandas as pd

ROOT = os.path.dirname(os.path.abspath(__file__))
EVAL = os.path.join(ROOT, "data/processed/evaluation")
CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

b = pd.read_csv(os.path.join(EVAL, "baselines.csv")).set_index("baseline")
bil_zw = float(b.loc["Bilinear interpolation", "RMSE_zw"])
clim_zw = float(b.loc["Climatology (training record)", "RMSE_zw"])

EDITS = [
    ("performance benchmarking, and benchmarking against the naive interpolation "
     "baseline defined in Section 3.7.3.",

     "performance benchmarking, and benchmarking against the training-period "
     "climatology, which Section 3.7.3 establishes as the only admissible skill "
     "reference here. The interpolation baseline defined in the same section is "
     "computed on this period too, but as a circularity diagnostic rather than as "
     "a competitor."),

    ("the first of them reproduces the target at a root-mean-square error of "
     "0.23 W/m², against the 19.08 W/m² root-mean-square error of the "
     "training-period climatology on the same record.",

     "the first of them reproduces the target at a root-mean-square error of "
     "%.2f W/m², against the %.2f W/m² root-mean-square error of the "
     "training-period climatology on the same record. Both figures are over "
     "Zimbabwe, the basis Table 4.1 and Section 4.6.1 also report."
     % (bil_zw, clim_zw)),
]


def is_field(r):
    x = r._r.xml
    return ("fldChar" in x) or ("instrText" in x)


def apply(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return True
    runs = par.runs
    full = "".join(r.text for r in runs)
    at = full.find(old)
    if at < 0:
        return False
    starts, pos = [], 0
    for r in runs:
        starts.append(pos)
        pos += len(r.text)
    f = max(i for i, st in enumerate(starts) if st <= at)
    l = max(i for i, st in enumerate(starts) if st < at + len(old))
    span = runs[f:l + 1]
    if any(is_field(r) for r in span):
        return False
    acc = "".join(r.text for r in span)
    k = acc.find(old)
    span[0].text = acc[:k] + new + acc[k + len(old):]
    for r in span[1:]:
        r.text = ""
    return True


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2
    d = docx.Document(DOC)
    x = d.element.xml
    before = (x.count("ZOTERO_ITEM"), x.count('w:fldCharType="begin"'))
    done = left = 0
    for old, new in EDITS:
        # Key idempotence on the tail: these replacements begin with the same
        # words as the text they replace, so a head match sees the original.
        if any(new[-60:] in p.text for p in d.paragraphs):
            print("  already applied: %s" % old[:52])
            continue
        par = next((p for p in d.paragraphs if old in p.text), None)
        if par is None:
            print("  NOT FOUND: %s" % old[:52])
            left += 1
            continue
        if apply(par, old, new):
            done += 1
            print("  fixed: %s" % old[:52])
        else:
            left += 1
            print("  SPANS A FIELD, skipped: %s" % old[:52])
    if not done:
        return 0 if left == 0 else 1
    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_basis.docx"))
    d.save(DOC)
    ax = docx.Document(DOC).element.xml
    after = (ax.count("ZOTERO_ITEM"), ax.count('w:fldCharType="begin"'))
    if after != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_basis.docx"), DOC)
        print("CITATION FIELDS CHANGED %s -> %s, restored" % (before, after))
        return 1
    print("%d fixed, %d outstanding; citation fields intact %s" % (done, left, before))
    return 0


if __name__ == "__main__":
    sys.exit(main())
