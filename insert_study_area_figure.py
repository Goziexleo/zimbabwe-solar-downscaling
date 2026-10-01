"""Add the study-area map to Section 3.2.1, and state the domain honestly.

Chapters 1 to 3 carried no figure at all. The map is placed where the domain is
defined, because that is where it does the most work: it shows the analysis box
against the national boundary, so the reader can see that the grid stops at 22 S
while the country reaches 22.42 S.

Section 3.2.1 also attributed the shortfall to ERA5 coverage, saying the retained
grid "is the extent over which the ERA5 predictor fields returned complete
coverage". ERA5 is global; the extent is what the download requested. The
sentence is corrected to say so, because an examiner who checks will find that
ERA5 covers 22.5 S perfectly well.
"""

import os
import shutil
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(ROOT, "figures/00_study_area.png")
CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

OLD = ("this being the extent over which the ERA5 predictor fields returned "
       "complete coverage")
NEW = ("this being the extent requested in the ERA5 download rather than a limit "
       "of the reanalysis, which is global. The national territory reaches 22.42° S, "
       "so the grid excludes a strip of southern Zimbabwe of roughly 51 cells, "
       "including Beitbridge; Section 4.9 records what this costs")

CAPTION = ("Figure 3.1. Zimbabwe's relief, provinces and the analysis domain. Elevation "
           "is SRTM at 0.1°. The dashed box is the retained analysis grid, 25–33° E and "
           "15–22° S; the national boundary extends south of it to 22.42° S, so "
           "Beitbridge and a strip of Matabeleland South and Masvingo fall outside "
           "the grid.")


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2
    if not os.path.exists(IMG):
        print("run make_study_area_map.py first")
        return 1

    d = docx.Document(DOC)
    x = d.element.xml
    before = (x.count("ZOTERO_ITEM"), x.count('w:fldCharType="begin"'),
              len(d.sections), x.count("<a:blip"))
    print("before:", before)

    par = next((p for p in d.paragraphs if OLD in p.text), None)
    if par is None:
        print("domain sentence already corrected")
    else:
        done = False
        for r in par.runs:
            if OLD in r.text:
                r.text = r.text.replace(OLD, NEW)
                done = True
                break
        if not done:
            runs = par.runs
            full = "".join(r.text for r in runs)
            at = full.find(OLD)
            starts, pos = [], 0
            for r in runs:
                starts.append(pos)
                pos += len(r.text)
            f = max(i for i, st in enumerate(starts) if st <= at)
            l = max(i for i, st in enumerate(starts) if st < at + len(OLD))
            span = runs[f:l + 1]
            if any(("fldChar" in r._r.xml or "instrText" in r._r.xml) for r in span):
                print("the sentence spans a citation field; left for the author")
            else:
                acc = "".join(r.text for r in span)
                k = acc.find(OLD)
                span[0].text = acc[:k] + NEW + acc[k + len(OLD):]
                for r in span[1:]:
                    r.text = ""
                done = True
        print("domain sentence corrected" if done else "NOT corrected")

    anchor = next(p for p in d.paragraphs if p.text.strip().startswith("3.2.2"))
    if any(p.text.strip().startswith("Figure 3.1.") for p in d.paragraphs):
        print("figure already present")
    else:
        fig = d.add_paragraph()
        fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fig.add_run().add_picture(IMG, width=Inches(6.0))
        cap = d.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(CAPTION)
        r.italic = True
        r.font.size = Pt(10)
        cap.paragraph_format.space_after = Pt(10)
        cap.paragraph_format.line_spacing = 1.0
        anchor._p.addprevious(fig._p)
        anchor._p.addprevious(cap._p)
        print("figure and caption inserted before Section 3.2.2")

    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_studyarea.docx"))
    d.save(DOC)
    y = docx.Document(DOC).element.xml
    after = (y.count("ZOTERO_ITEM"), y.count('w:fldCharType="begin"'),
             len(docx.Document(DOC).sections), y.count("<a:blip"))
    print("after: ", after)
    if after[:3] != before[:3] or after[3] != before[3] + 1:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_studyarea.docx"), DOC)
        print("UNEXPECTED CHANGE, restored")
        return 1
    print("citation fields and sections intact; one image added")
    return 0


if __name__ == "__main__":
    sys.exit(main())
