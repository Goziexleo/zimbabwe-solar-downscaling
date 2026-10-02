"""Replace the clear-sky disclosure with the result, now that the basis is fixed.

Section 3.4.3 originally said the 1.1 clip "does not bind anywhere in the record"
and that "the number of flagged ERA5 values is exactly zero", offered as evidence
that ERA5 never exceeds its own clear-sky ceiling. That was an artefact: the
denominator was a daytime mean and the numerator a 24-hour mean, so the index
could not exceed 0.6323 and the test could not bind.

compute_finegrid_clearsky_ghi.py now integrates the full 24 hours, the target has
been rebuilt and the models refitted. The index averages 0.862 and the clip binds
on 441 of 1,794,312 training values and 385 of 966,168 validation values. The
paragraph is rewritten to report that rather than the artefact or the caveat.
"""

import os
import re
import shutil
import sys

import docx

CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

ANCHOR = "The constraint does not bind anywhere in the record"
END = "it is recorded as outstanding in Section 4.9."

NEW = (
    "The constraint binds, and rarely. Across 1,794,312 training and 966,168 "
    "validation values the clear-sky index averages 0.862, spans 0.661 to 1.014 "
    "between the fifth and ninety-fifth percentiles, and is truncated at the upper "
    "bound for 441 training and 385 validation values, 0.030 per cent of the "
    "record. Those are months in which the bicubically interpolated ERA5 field "
    "exceeds the modelled clear-sky ceiling by more than ten per cent, which is "
    "the behaviour the bound exists to catch. An earlier version of this section "
    "reported the number of flagged values as exactly zero, with a maximum index "
    "of 0.6323. That was an artefact of the denominator: the clear-sky ceiling was "
    "averaged over thirteen samples spanning 06:00 to 18:00, giving a daytime mean "
    "of 502.7 W/m², while the target it normalises is a 24-hour mean of 237.3 "
    "W/m² formed from ERA5 daily accumulations. The ratio of the two conventions "
    "is close to two, so the index could not approach unity and the bound could not "
    "bind; its passing was therefore not evidence about ERA5. The timestamps were "
    "also naive and read as UTC, placing that window at 08:00 to 20:00 local. The "
    "ceiling is now integrated over the whole day, which removes both faults, and "
    "every model, projection and suitability layer downstream has been refitted "
    "against the rebuilt target.")


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2
    d = docx.Document(DOC)
    x = d.element.xml
    before = (x.count("ZOTERO_ITEM"), x.count('w:fldCharType="begin"'))
    print("before:", before)

    par = next((p for p in d.paragraphs if ANCHOR in p.text), None)
    if par is None:
        print("anchor not found; already rewritten?")
        return 0
    if "The constraint binds, and rarely" in par.text:
        print("already applied")
        return 0

    full = "".join(r.text for r in par.runs)
    a = full.find(ANCHOR)
    b = full.find(END)
    if b < 0:
        print("could not find the end of the passage; left alone")
        return 1
    b += len(END)

    starts, pos = [], 0
    for r in par.runs:
        starts.append(pos)
        pos += len(r.text)
    f = max(i for i, st in enumerate(starts) if st <= a)
    l = max(i for i, st in enumerate(starts) if st < b)
    span = par.runs[f:l + 1]
    if any(("fldChar" in r._r.xml or "instrText" in r._r.xml) for r in span):
        print("the passage spans a citation field; left for the author")
        return 1

    acc = "".join(r.text for r in span)
    k = acc.find(ANCHOR)
    j = acc.find(END) + len(END)
    span[0].text = acc[:k] + NEW + acc[j:]
    for r in span[1:]:
        r.text = ""

    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_qcclaim.docx"))
    d.save(DOC)
    y = docx.Document(DOC).element.xml
    after = (y.count("ZOTERO_ITEM"), y.count('w:fldCharType="begin"'))
    if after != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_qcclaim.docx"), DOC)
        print("CITATION FIELDS CHANGED, restored")
        return 1
    print("rewritten; citation fields intact")
    return 0


if __name__ == "__main__":
    sys.exit(main())
