"""Add the fourth use to Appendix B: summarising literature, and what it cost.

The declaration listed three uses - language, code generation, debugging - and
stated that the tool produced no number reported in the dissertation. Both are
true, but incomplete: the tool was also used to summarise sources, and an audit
against the papers themselves found attributions several of them do not support.
Buster et al. (2024) was described as training on ERA5 when it used ERA5 as the
low-resolution input and NSRDB and WTK as targets; Rampal et al. (2024) was
described as a benchmarking study when it is a review; Lin et al. (2023) was cited
for Zimbabwean irradiance although it covers East Asia.

A declaration that omits the use which actually produced errors is worse than no
declaration, because an examiner who finds the errors will conclude the
declaration was drafted to look complete rather than to be complete. The honest
version names the use, says what it cost, and says what was done about it.

Also adds Grammarly, whose metadata the file carries.
"""

import os
import shutil
import sys

import docx

CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

OLD_USES = "The tool was used for three purposes."
NEW_USES = "The tool was used for four purposes."

OLD_TAIL = ("a set of figures whose plotted labels had fallen out of step with the "
            "analysis behind them.")
NEW_TAIL = (
    "a set of figures whose plotted labels had fallen out of step with the "
    "analysis behind them. Fourth, literature: summarising the content of cited "
    "papers.")

NEW_PARA = (
    "The fourth use proved the least reliable and is declared with the detail it "
    "warrants. Summaries of cited work produced by the tool were in several cases "
    "not supported by the papers themselves. Buster et al. (2024) was described as "
    "having trained against ERA5, when ERA5 was its low-resolution input and the "
    "National Solar Radiation Database and Wind Integration National Dataset were "
    "its high-resolution targets; Rampal et al. (2024) was described as a "
    "benchmarking study when it is a review; and Lin et al. (2023), a dataset for "
    "East Asia, was cited for Zimbabwean irradiance. Every citation was "
    "subsequently checked against the source, the affected passages were rewritten, "
    "and the audit is recorded in the repository. The errors were the author's to "
    "catch and the responsibility for them is the author's; they are set out here "
    "because a declaration that omitted the use which caused them would not be a "
    "declaration.")

GRAMMARLY = (
    "Grammarly was also used for spelling and grammar checking, and the working "
    "file carries its metadata.")


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2

    d = docx.Document(DOC)
    x = d.element.xml
    before = (x.count("ZOTERO_ITEM"), x.count('w:fldCharType="begin"'))

    uses = next((p for p in d.paragraphs if OLD_USES in p.text), None)
    if uses is None:
        print("the three-purposes sentence is not present; already updated?")
    else:
        for r in uses.runs:
            if OLD_USES in r.text:
                r.text = r.text.replace(OLD_USES, NEW_USES)
            if OLD_TAIL in r.text:
                r.text = r.text.replace(OLD_TAIL, NEW_TAIL)
        print("declared a fourth use")

    anchor = next(p for p in d.paragraphs
                  if p.text.strip().startswith("The tool was not used to generate"))
    if not any(p.text.strip().startswith("The fourth use proved") for p in d.paragraphs):
        new = d.add_paragraph(NEW_PARA, style=anchor.style)
        anchor._p.addprevious(new._p)
        new.alignment = anchor.alignment
        pf, src = new.paragraph_format, anchor.paragraph_format
        pf.space_before, pf.space_after = src.space_before, src.space_after
        pf.line_spacing, pf.first_line_indent = src.line_spacing, src.first_line_indent
        if anchor.runs:
            for r in new.runs:
                r.font.name = anchor.runs[0].font.name
                r.font.size = anchor.runs[0].font.size
        print("added the literature-summary paragraph")

    last = next(p for p in d.paragraphs
                if p.text.strip().startswith("Responsibility for the content"))
    if GRAMMARLY not in last.text:
        last.runs[-1].text = last.runs[-1].text.rstrip() + " " + GRAMMARLY
        print("declared Grammarly")

    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_appendixB.docx"))
    d.save(DOC)
    y = docx.Document(DOC).element.xml
    after = (y.count("ZOTERO_ITEM"), y.count('w:fldCharType="begin"'))
    if after != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_appendixB.docx"), DOC)
        print("CITATION FIELDS CHANGED, restored")
        return 1
    print("citation fields intact (%d)" % after[0])
    return 0


if __name__ == "__main__":
    sys.exit(main())
