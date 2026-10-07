"""Bring Table 3.4 and its two sentences onto the deployed configuration.

The sixth critique's third item. The ablation's configurations were one-off
refits that no script re-executed, so when the pixel-wise hyperparameters changed
the table went stale, and a limitations sentence ended up reading "5.54 against
8.88" — the first figure from the old configuration on the full analysis box, the
second from the new one over Zimbabwe. Two configurations and two bases in one
comparison.

compute_rsds_ablation.py now reproduces B and C at the deployed hyperparameters.
Configuration A is not reproducible: it required the one-month predictor
misalignment, which was a defect in the feature builder rather than a switch, so
its row is kept and labelled as a measurement under the original configuration.
Skill is against the same climatology Table 4.1 uses.

    python fix_ablation_table.py
"""

import os
import shutil
import sys

import docx
import pandas as pd

ROOT = os.path.dirname(os.path.abspath(__file__))
EVAL = os.path.join(ROOT, "data/processed/evaluation")
CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

abl = pd.read_csv(os.path.join(EVAL, "rsds_ablation.csv"))
t33 = pd.read_csv(os.path.join(EVAL, "table_3_3.csv")).set_index("model")


def row(prefix):
    r = abl[abl.configuration.str.startswith(prefix)]
    assert not r.empty, prefix
    return r.iloc[0]


B, C = row("B."), row("C.")
clim = float(t33["climatology_rmse"].iloc[0])
clim_zw = float(t33["climatology_rmse_zw"].iloc[0])

CELLS = {
    "B. Alignment corrected, rsds retained":
        ["%.2f" % B.RMSE_fullbox, "%.4f" % B.pearson_r,
         "%.4f" % (1.0 - B.RMSE_fullbox / clim)],
    "C. Alignment corrected, rsds excluded (deployed)":
        ["%.2f" % C.RMSE_fullbox, "%.4f" % C.pearson_r,
         "%.4f" % (1.0 - C.RMSE_fullbox / clim)],
    "A. Misaligned rsds, six predictors (as originally reported)":
        None,   # not reproducible; label it instead
}

EDITS = [
    ("A. Misaligned rsds, six predictors (as originally reported)",
     "A. Misaligned rsds, six predictors (original configuration, not re-measured)"),

    ("raising the error from 5.54 to 9.03 W/m², both over the full analysis box on "
     "which the ablation was run",
     "raising the error from %.2f to %.2f W/m², both over the full analysis box on "
     "which the ablation was run and both at the deployed hyperparameters"
     % (B.RMSE_fullbox, C.RMSE_fullbox)),

    ("with the field retained the same model reaches 5.54 W/m² against 8.88 without "
     "it",
     "with the field retained the same model reaches %.2f W/m² against %.2f without "
     "it, both over Zimbabwe" % (B.RMSE_zw, C.RMSE_zw)),
]


def is_field(r):
    x = r._r.xml
    return ("fldChar" in x) or ("instrText" in x)


def apply(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return True
    runs = par.runs
    full = "".join(r.text for r in runs)
    at = full.find(old)
    if at < 0:
        return False
    starts, pos = [], 0
    for r in runs:
        starts.append(pos)
        pos += len(r.text)
    f = max(i for i, st in enumerate(starts) if st <= at)
    l = max(i for i, st in enumerate(starts) if st < at + len(old))
    span = runs[f:l + 1]
    if any(is_field(r) for r in span):
        return False
    acc = "".join(r.text for r in span)
    k = acc.find(old)
    span[0].text = acc[:k] + new + acc[k + len(old):]
    for r in span[1:]:
        r.text = ""
    return True


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2
    d = docx.Document(DOC)
    before = (d.element.xml.count("ZOTERO_ITEM"),
              d.element.xml.count('w:fldCharType="begin"'))
    changed = 0

    tbl = next((t for t in d.tables
                if any("Alignment corrected" in c.text for r in t.rows for c in r.cells)),
               None)
    assert tbl is not None, "Table 3.4 not found"
    for r in tbl.rows[1:]:
        label = r.cells[0].text.strip()
        vals = CELLS.get(label)
        if not vals:
            continue
        for cell, v in zip(r.cells[1:], vals):
            if cell.text.strip() != v:
                cell.paragraphs[0].runs[0].text = v
                changed += 1
    print("  %d table cells updated" % changed)

    cells = [p for t in d.tables for row_ in t.rows for c in row_.cells
             for p in c.paragraphs]
    for old, new in EDITS:
        par = next((p for p in list(d.paragraphs) + cells if old in p.text), None)
        if par is None:
            print("  already applied or absent: %s" % old[:52])
            continue
        if apply(par, old, new):
            changed += 1
            print("  fixed: %s" % old[:52])
        else:
            print("  SPANS A FIELD, skipped: %s" % old[:52])

    if not changed:
        return 0
    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_ablation34.docx"))
    d.save(DOC)
    ax = docx.Document(DOC).element.xml
    after = (ax.count("ZOTERO_ITEM"), ax.count('w:fldCharType="begin"'))
    if after != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_ablation34.docx"), DOC)
        print("CITATION FIELDS CHANGED %s -> %s, restored" % (before, after))
        return 1
    print("%d changes; citation fields intact %s" % (changed, after))
    return 0


if __name__ == "__main__":
    sys.exit(main())
