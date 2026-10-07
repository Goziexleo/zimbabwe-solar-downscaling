"""Cite Taylor (2001) and resolve the orphaned Cannon (2018) entry.

Two citation faults the v7 re-check raised:

  - Section 3.7.2 describes the Taylor diagram and computes Taylor statistics
    without citing Taylor (2001). The entry is prepared in
    reference_additions.ris with a CrossRef-verified DOI but is not in the
    document's bibliography and nothing in the body points at it.
  - Cannon (2018), the multivariate quantile-mapping paper, is listed in the
    bibliography and cited nowhere. An uncited entry is a padding signal.

Cannon (2018) is cited where it belongs rather than deleted: Section 3.4.4
corrects each predictor independently, which does not preserve the dependence
between them, and the N-dimensional transform is the method that would. Saying
so turns an orphaned entry into a stated limitation.

Also corrects "each deployed model" in Section 3.4.4's pointer to Table C.9.
Only XGBoost is deployed; the Random Forest is a comparison.

Both citations are inserted as plain text. insert_missing_references.py then
adds any that the bibliography lacks, and a Zotero refresh in Word replaces the
plain-text form with a field.

    python fix_method_citations_v7.py
"""

import os
import shutil
import sys

import docx

CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

EDITS = [
    ("First, the Taylor diagram framework is used to simultaneously display the "
     "spatial correlation, centred RMSE, and standard deviation ratio",
     "First, the Taylor diagram framework (Taylor, 2001) is used to "
     "simultaneously display the spatial correlation, centred RMSE, and "
     "standard deviation ratio"),

    ("Quantile mapping (QM) bias correction was therefore applied to each CMIP6 "
     "predictor variable independently, before those fields are presented to "
     "the already-trained models at the projection stage, rather than before "
     "training itself.",
     "Quantile mapping (QM) bias correction was therefore applied to each CMIP6 "
     "predictor variable independently, before those fields are presented to "
     "the already-trained models at the projection stage, rather than before "
     "training itself. Correcting each variable on its own leaves the "
     "dependence between them as the GCM produced it, so a model that pairs "
     "cloud fraction with temperature in a way ERA5 does not is corrected "
     "variable by variable and not jointly; the N-dimensional transform of "
     "Cannon (2018) is the method that would address this, and it was not "
     "applied here. Section 4.7.1's transfer test is the measurement of what "
     "that leaves behind."),

    ("Appendix C, Table C.9, applies each deployed model to every driving "
     "model's bias-corrected predictors",
     "Appendix C, Table C.9, applies each of the two pixel-wise models to every "
     "driving model's bias-corrected predictors"),
]


def is_field(r):
    x = r._r.xml
    return ("fldChar" in x) or ("instrText" in x)


def apply(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return True
    runs = par.runs
    full = "".join(r.text for r in runs)
    at = full.find(old)
    if at < 0:
        return False
    starts, pos = [], 0
    for r in runs:
        starts.append(pos)
        pos += len(r.text)
    f = max(i for i, st in enumerate(starts) if st <= at)
    l = max(i for i, st in enumerate(starts) if st < at + len(old))
    span = runs[f:l + 1]
    if any(is_field(r) for r in span):
        return False
    acc = "".join(r.text for r in span)
    k = acc.find(old)
    span[0].text = acc[:k] + new + acc[k + len(old):]
    for r in span[1:]:
        r.text = ""
    return True


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2
    d = docx.Document(DOC)
    before = (d.element.xml.count("ZOTERO_ITEM"),
              d.element.xml.count('w:fldCharType="begin"'), len(d.sections))
    done = left = 0
    for old, new in EDITS:
        i = 0
        while i < min(len(old), len(new)) and old[i] == new[i]:
            i += 1
        key = new[i:i + 70].strip()
        if key and any(key in p.text for p in d.paragraphs):
            print("  already applied: %s" % old[:50])
            continue
        par = next((p for p in d.paragraphs if old in p.text), None)
        if par is None:
            if any(new in p.text for p in d.paragraphs):
                print("  already applied: %s" % old[:50])
            else:
                left += 1
                print("  NOT FOUND: %s" % old[:50])
            continue
        if apply(par, old, new):
            done += 1
            print("  cited: %s" % key[:60])
        else:
            left += 1
            print("  SPANS A FIELD, skipped: %s" % old[:50])
    if not done:
        print("nothing to change")
        return 0 if left == 0 else 1
    bak = DOC.replace(".docx", "_BACKUP_pre_citations_v7.docx")
    shutil.copy(DOC, bak)
    d.save(DOC)
    dd = docx.Document(DOC)
    ax = dd.element.xml
    after = (ax.count("ZOTERO_ITEM"), ax.count('w:fldCharType="begin"'),
             len(dd.sections))
    if after != before:
        shutil.copy(bak, DOC)
        print("STRUCTURE CHANGED %s -> %s, restored" % (before, after))
        return 1
    print("%d cited, %d outstanding; citations/fields/sections intact %s"
          % (done, left, after))
    return 0


if __name__ == "__main__":
    sys.exit(main())
