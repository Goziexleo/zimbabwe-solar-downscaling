"""Bring Chapter 3's hand-written figures into line with the rebuilt analysis.

Chapters 4 and 5 regenerate from the evaluation CSVs, so they followed the
clear-sky correction and the U-Net reseed on their own. Chapter 3 is written by
hand and did not, which left it contradicting Chapter 4 on a claim rather than
only on digits.

The substantive change is in Section 3.8.4. The Random Forest's centred-error
advantage over XGBoost used to clear a BCa interval, -0.171 at [-0.360, -0.041],
and the section concluded that the composite criterion was supported on both of
its legs. On the rebuilt analysis that difference is -0.133 at [-0.344, +0.009],
which contains zero. The criterion now stands on one leg, the systematic offset,
where the Random Forest's advantage is both established and larger than before.
That makes the deployment argument simpler, not weaker, and the section is
rewritten to say so.

The remaining replacements are figures: validation errors, centred errors, the
paired intervals and the scenario separation.
"""

import os
import shutil
import sys

import docx

CH = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-Personal(2)/UNI ZIM/PROJECT CHAPTERS")
DOC = os.path.join(CH, "CR_Madukwe_Dissertation.docx")
LOCK = os.path.join(CH, "~$_Madukwe_Dissertation.docx")

OLD_BLOCK = (
    "A percentile interval places its centred-error advantage at -0.145 W/m² with "
    "a 95 per cent interval of [-0.305, +0.010], which contains zero; but that "
    "interval is not appropriate here, because each bootstrap replicate recomputes "
    "the time-mean field from resampled years and so acquires sampling noise in "
    "quadrature, which biases replicate differences toward zero. A bias-corrected "
    "and accelerated interval, which is the standard remedy, gives -0.171 W/m² at "
    "[-0.360, -0.041] and excludes zero. The Random Forest's centred-error "
    "advantage is therefore established. Its spatial-correlation advantage is not, "
    "at +0.0015 with a bias-corrected interval of [-0.001, +0.004]. What follows is "
    "that the criterion supports the Random Forest on both of its legs, not on one "
    "and not on neither, and that the deployment argument cannot rest on denying "
    "any part of it. It does not need to. The advantage is 0.171 W/m² on a "
    "time-mean field, while the aggregate deficit the criterion was introduced to "
    "outweigh is six times larger and is itself established, at +1.058 W/m² with a "
    "bias-corrected interval of [+0.598, +1.768], and XGBoost returns the lower "
    "error in all four folds of a rolling-origin evaluation over expanding windows."
)

NEW_BLOCK = (
    "A percentile interval places its centred-error advantage at -0.133 W/m² with "
    "a 95 per cent interval of [-0.289, +0.067], which contains zero. A "
    "bias-corrected and accelerated interval, which corrects the downward bias that "
    "arises because each bootstrap replicate recomputes the time-mean field from "
    "resampled years and so acquires sampling noise in quadrature, gives -0.133 "
    "W/m² at [-0.344, +0.009]. That also contains zero, and the Random Forest's "
    "centred-error advantage is therefore not established. Nor is its "
    "spatial-correlation advantage, at +0.0010 with a bias-corrected interval of "
    "[-0.002, +0.003]. What follows is that the criterion stands on one of its two "
    "legs rather than both: the systematic offset, where the Random Forest is "
    "established better by 1.194 W/m², and not the spatial structure of the error, "
    "where the two cannot be separated. An earlier state of this analysis found the "
    "second leg established as well; the rebuilt target of Section 3.4.3 moved it "
    "across the threshold, and the weaker claim is the one the evidence now "
    "supports. The deployment argument is correspondingly simpler. The aggregate "
    "deficit the criterion was introduced to outweigh is itself established, at "
    "+1.183 W/m² with a bias-corrected interval of [+0.694, +2.069], and XGBoost "
    "returns the lower error in all four folds of a rolling-origin evaluation over "
    "expanding windows."
)

EDITS = [
    (OLD_BLOCK, NEW_BLOCK),
    # Section 3.8.4, first leg
    ("established better, by 0.953 W/m² with a bias-corrected interval of "
     "[-1.499, -0.475]",
     "established better, by 1.194 W/m² with a bias-corrected interval of "
     "[-1.764, -0.484]"),
    # scenario separation
    ("from 0.465 W/m² in the near term to 0.167 in the long term",
     "from 0.464 W/m² in the near term to 0.169 in the long term"),
    ("only 71.7 per cent of grid cells order", "only 71.5 per cent of grid cells order"),
    # Section 3.7.2 centred errors
    ("at 0.571 W/m² for the Random Forest, 0.742 for XGBoost, 3.076 for the CNN "
     "and 3.723 for the U-Net",
     "at 0.554 W/m² for the Random Forest, 0.687 for XGBoost, 3.155 for the CNN "
     "and 4.341 for the U-Net"),
    ("where the centred errors differ by a factor of about five",
     "where the centred errors differ by a factor of about six"),
    ("+0.217 W/m² against the CNN", "+0.181 W/m² against the CNN"),
    # the cost of deploying XGBoost
    ("Its mean bias is +1.20 W/m² against the Random Forest's +0.25",
     "Its mean bias is +1.42 W/m² against the Random Forest's +0.22"),
    ("the paired difference in mean bias is -0.953 W/m² with a bias-corrected "
     "interval of [-1.499, -0.475]",
     "the paired difference in mean bias is -1.194 W/m² with a bias-corrected "
     "interval of [-1.764, -0.484]"),
    ("For scale, +1.20 is roughly an eighth of a downscaling uncertainty of order "
     "10 W/m²", "For scale, +1.42 is roughly a seventh of a downscaling "
     "uncertainty of order 10 W/m²"),
    # Section 3.6.4 and 3.7.4
    ("the cost of removing it was 9.11 against the 9.24 reported here",
     "the cost of removing it was 9.11 against the 9.03 reported here"),
    ("raising the error from 5.54 to 9.24 W/m²",
     "raising the error from 5.54 to 9.03 W/m²"),
    ("the lower aggregate error, at 9.24 W/m² against 10.30",
     "the lower aggregate error, at 9.03 W/m² against 10.22"),
]


def is_field(r):
    x = r._r.xml
    return ("fldChar" in x) or ("instrText" in x)


def apply(par, old, new):
    for r in par.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return True
    runs = par.runs
    full = "".join(r.text for r in runs)
    at = full.find(old)
    if at < 0:
        return False
    starts, pos = [], 0
    for r in runs:
        starts.append(pos)
        pos += len(r.text)
    f = max(i for i, st in enumerate(starts) if st <= at)
    l = max(i for i, st in enumerate(starts) if st < at + len(old))
    span = runs[f:l + 1]
    if any(is_field(r) for r in span):
        return False
    acc = "".join(r.text for r in span)
    k = acc.find(old)
    span[0].text = acc[:k] + new + acc[k + len(old):]
    for r in span[1:]:
        r.text = ""
    return True


def main():
    if os.path.exists(LOCK):
        print("The dissertation is open in Word. Close it first.")
        return 2
    d = docx.Document(DOC)
    x = d.element.xml
    before = (x.count("ZOTERO_ITEM"), x.count('w:fldCharType="begin"'))
    print("before:", before)

    done = left = 0
    for old, new in EDITS:
        par = next((p for p in d.paragraphs if old in p.text), None)
        if par is None:
            if any(new in p.text for p in d.paragraphs):
                print("  already applied: %s" % old[:52])
            else:
                print("  NOT FOUND: %s" % old[:52])
                left += 1
            continue
        if apply(par, old, new):
            done += 1
            print("  fixed: %s" % old[:52])
        else:
            left += 1
            print("  SPANS A FIELD: %s" % old[:52])

    if not done:
        print("nothing to do")
        return 0 if left == 0 else 1
    shutil.copy(DOC, DOC.replace(".docx", "_BACKUP_pre_ch3.docx"))
    d.save(DOC)
    y = docx.Document(DOC).element.xml
    after = (y.count("ZOTERO_ITEM"), y.count('w:fldCharType="begin"'))
    if after != before:
        shutil.copy(DOC.replace(".docx", "_BACKUP_pre_ch3.docx"), DOC)
        print("CITATION FIELDS CHANGED, restored")
        return 1
    print("%d fixed, %d outstanding; citation fields intact" % (done, left))
    return 0


if __name__ == "__main__":
    sys.exit(main())
